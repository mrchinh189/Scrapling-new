import csv
from pathlib import Path

from nvl_scraper.models import PriceRecord

KEY_FIELDS = ("source", "material", "series", "price_date")


def _key(row: dict) -> tuple:
    return tuple(row[k] for k in KEY_FIELDS)


def load_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def merge_into_csv(path: Path, records: list[PriceRecord]) -> tuple[list[dict], int]:
    """Upsert records into the history CSV (one row per source/material/series/date).

    Returns all rows (sorted) and the number of new rows added.
    """
    rows = {_key(r): r for r in load_csv(path)}
    before = len(rows)
    for rec in records:
        row = {k: ("" if v is None else v) for k, v in rec.to_row().items()}
        rows[_key(row)] = row  # newer scrape wins (prices can be revised)
    merged = sorted(rows.values(), key=lambda r: (r["price_date"], r["source"], r["material"], r["series"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel opens Vietnamese/Chinese text correctly
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=PriceRecord.columns())
        writer.writeheader()
        writer.writerows(merged)
    return merged, len(rows) - before


def latest_per_series(rows: list[dict]) -> list[dict]:
    latest: dict[tuple, dict] = {}
    for row in rows:
        k = (row["source"], row["material"], row["series"])
        if k not in latest or row["price_date"] >= latest[k]["price_date"]:
            latest[k] = row
    return sorted(latest.values(), key=lambda r: (r["material"], r["series"]))


def write_excel(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    columns = PriceRecord.columns()
    numeric = {"price", "change", "change_pct"}
    wb = Workbook()
    sheets = (("Mới nhất", latest_per_series(rows)), ("Lịch sử", rows))
    for i, (title, data) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(columns)
        for row in data:
            ws.append([_cell(row.get(c), c in numeric) for c in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, col in enumerate(columns, 1):
            width = max([len(col)] + [len(str(r.get(col, ""))) for r in data[:500]])
            ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _cell(value, numeric: bool):
    if value in (None, ""):
        return None
    if numeric:
        try:
            return float(value)
        except ValueError:
            return value
    return value
