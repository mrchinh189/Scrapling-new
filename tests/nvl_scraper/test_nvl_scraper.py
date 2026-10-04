import csv
import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest

pytest.importorskip("openpyxl")

from scrapling.parser import Selector

from nvl_scraper.sources import InvestingSource, LmeSource, MpocSource, Ppi100Source
from nvl_scraper.storage import latest_per_series, merge_into_csv, write_excel
from nvl_scraper.utils import infer_year, parse_number

FIXTURES = Path(__file__).parent / "fixtures"


def page(name: str) -> Selector:
    return Selector((FIXTURES / name).read_text(encoding="utf-8"))


def test_parse_number():
    assert parse_number("1,121.25") == 1121.25
    assert parse_number("(-0.47%)") == -0.47
    assert parse_number("-") is None
    assert parse_number(None) is None


def test_infer_year_rolls_back_over_new_year():
    assert infer_year(12, 30, today=date(2027, 1, 3)) == date(2026, 12, 30)
    assert infer_year(10, 3, today=date(2026, 10, 4)) == date(2026, 10, 3)


def test_investing():
    src = InvestingSource("investing_cpo", "https://example.com", "CPO")
    [rec] = src.parse(page("investing.html"))
    assert (rec.price, rec.change, rec.change_pct) == (1121.25, -5.25, -0.47)
    assert rec.price_date == date(2026, 10, 2)
    assert rec.series == "CPOc1"


def test_100ppi_skips_unpublished_today():
    src = Ppi100Source("100ppi_h2so4", "https://example.com", "H2SO4")
    recs = src.parse(page("ppi100.html"), today=date(2026, 10, 4))
    assert [r.price_date for r in recs] == [date(2026, 10, 3), date(2026, 10, 2), date(2026, 9, 30), date(2026, 9, 12)]
    assert recs[0].price == 1710.0
    assert recs[2].change_pct == -1.58


def test_mpoc_both_tables():
    recs = MpocSource().parse(page("mpoc.html"))
    settlement = [r for r in recs if r.series == "Bursa settlement"]
    assert [(r.price_date, r.price) for r in settlement] == [(date(2026, 9, 30), 4610), (date(2026, 10, 1), 4554)]
    monthly = [r for r in recs if r.series.startswith("Monthly")]
    assert len(monthly) == 6  # 2 months x CPO/SBO/SFO, premiums ignored
    sfo_apr = next(r for r in monthly if "SFO" in r.material and r.price_date == date(2026, 4, 1))
    assert sfo_apr.price == 1295 and sfo_apr.currency == "USD"


def test_lme_html_fallback():
    src = LmeSource("lme_zinc", "https://example.com", "Zinc")
    recs = src.parse(page("lme_table.html"))
    assert {r.series: r.price for r in recs} == {
        "Cash Bid": 2951.5,
        "Cash Offer": 2952.0,
        "3-months Bid": 2970.0,
        "3-months Offer": 2971.0,
    }


def test_lme_api_json():
    payload = {
        "DateOfData": "2026-10-02T00:00:00",
        "Rows": [
            {"Label": "Cash", "Values": ["2,951.50", "2,952.00"]},
            {"Label": "3-months", "Values": ["2,970.00", "2,971.00"]},
            {"Label": "Dec 27", "Values": ["1", "2"]},
        ],
    }
    pg = page("lme_table.html")
    pg.captured_xhr = [SimpleNamespace(url="x", body=json.dumps(payload).encode())] * 2
    recs = LmeSource("lme_zinc", "https://example.com", "Zinc").parse(pg)
    assert len(recs) == 4  # duplicated API calls are de-duplicated
    assert all(r.price_date == date(2026, 10, 2) for r in recs)


def test_storage_upsert_and_excel(tmp_path):
    mpoc = MpocSource().parse(page("mpoc.html"))
    csv_path = tmp_path / "nvl_prices.csv"
    rows, added = merge_into_csv(csv_path, mpoc)
    assert added == len(mpoc)
    rows, added = merge_into_csv(csv_path, mpoc)  # re-run same day: no duplicates
    assert added == 0
    with csv_path.open(encoding="utf-8-sig") as fh:
        assert len(list(csv.DictReader(fh))) == len(mpoc)

    latest = latest_per_series(rows)
    cpo = next(r for r in latest if r["series"] == "Bursa settlement")
    assert cpo["price_date"] == "2026-10-01"

    xlsx = tmp_path / "nvl_prices.xlsx"
    write_excel(xlsx, rows)
    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Mới nhất", "Lịch sử"]
    assert wb["Lịch sử"].max_row == len(mpoc) + 1
